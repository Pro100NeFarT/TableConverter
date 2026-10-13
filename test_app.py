"""Автотест Excel-потоку через GUI-логіку (без кліків мишею). Файл генерується самим тестом."""
import csv
import io
import os
import tempfile

import openpyxl
from openpyxl.styles import Font

import converter
import sources
from app import App

SRC = os.path.join(tempfile.gettempdir(), "tc_test.xlsx")
OUT = os.path.join(tempfile.gettempdir(), "tc_test_excel.csv")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Товари"
ws.append([None, "КОД", "Назва", "Склад", "Опис"])
ws.append([None, "НФ-1", 'Мило "Ерл грей" 250 мл', "Вода, гліцерин; пантенол", "Рядок 1\nрядок 2"])
ws.append([None, "НФ-2", "Крем", "Вода | олія", None])
ws.append([None, "НФ-3", "Гель", "Вода: вода", "Ок"])
ws["B200"].font = Font(bold=True)  # розтягує використаний діапазон: порожній хвіст
wb.create_sheet("Інше").append(["a", "b"])
wb.save(SRC)

app = App()
app.file_path = SRC
app.reload()
assert [t.name for t in app.tables] == ["Товари", "Інше"], [t.name for t in app.tables]
app.tables_list.selection_clear(0, "end")
app.tables_list.selection_set(0)

app.analyze()
delim = app.get_selected_delimiter()
print("Обраний роздільник:", repr(delim))
assert delim == "\t"  # ';' , '|' , ':' зустрічаються в даних
assert app._cells_with_newline == 1 and app._cells_with_quote == 1
assert sum(app._trailing_empty.values()) > 150, app._trailing_empty

converter.convert_rows_to_csv(app.tables[0].rows, delim, OUT)
with open(OUT, "r", encoding="utf-8-sig", newline="") as f:
    content = f.read()
os.remove(OUT)
os.remove(SRC)

assert "\n" not in content.replace("\r\n", "") and '"' not in content
rows = list(csv.reader(io.StringIO(content), delimiter=delim))
assert len(rows) == 4 and content.count("\r\n") == 4, len(rows)
assert {len(r) for r in rows} == {5}
assert rows[1][2] == "Мило 'Ерл грей' 250 мл" and rows[1][4] == "Рядок 1 рядок 2"
assert rows[1][3] == "Вода, гліцерин; пантенол"
print("Excel: 4 рядки x 5 колонок, без переносів, лапок і порожнього хвоста.")

app.destroy()
print("ВСІ ПЕРЕВІРКИ EXCEL ПРОЙДЕНО")
