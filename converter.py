"""
Логіка аналізу та запису CSV/Excel (без GUI, незалежна від джерела: Excel/PDF).
Працює зі списками рядків (rows) і sources.Table.
"""
import csv
import io
import re
from collections import Counter

# Кандидати на роздільник: (символ, назва, пріоритет показу)
DELIMITER_CANDIDATES = [
    (";", "крапка з комою", 1),
    ("\t", "табуляція (Tab)", 2),
    ("|", "вертикальна риска (pipe)", 3),
    ("^", "циркумфлекс", 4),
    ("~", "тильда", 5),
    ("¦", "розірвана риска (broken bar)", 6),
    ("§", "параграф", 7),
    ("#", "решітка", 8),
    ("@", "ет (at)", 9),
    ("¶", "пілкроу", 10),
    ("‡", "подвійний хрестик", 11),
    ("†", "хрестик", 12),
    ("¤", "валютний символ", 13),
    ("№", "номер", 14),
    ("«", "лапка-ялинка відкр.", 15),
    ("»", "лапка-ялинка закр.", 16),
    (",", "кома", 17),
    (":", "двокрапка", 18),
    ("\x1f", "Unit Separator (керівний символ, невидимий)", 19),
]
_PRIORITY = {ch: prio for ch, _, prio in DELIMITER_CANDIDATES}

_SPACES = "   "
_NUM_SPACED = re.compile(r"[+-]?\d{1,3}(?:[   ]\d{3})+(?:[.,]\d+)?")
_NUM_DECIMAL_COMMA = re.compile(r"[+-]?\d+,\d+")


def _is_empty_row(row):
    return all(v is None or str(v).strip() == "" for v in row)


def get_data_rows(rows, trim_trailing_empty=True):
    """Рядки таблиці (колонки не змінюються). Опційно відкидає порожні рядки в кінці."""
    if not trim_trailing_empty:
        return rows
    last_idx = -1
    for i, row in enumerate(rows):
        if not _is_empty_row(row):
            last_idx = i
    return rows[: last_idx + 1]


def count_trailing_empty_rows(rows):
    return len(rows) - len(get_data_rows(rows, trim_trailing_empty=True))


def analyze_characters(tables):
    """Частота символів по обраних таблицях.
    Повертає (Counter, клітинок, клітинок_з_переносом_рядка, клітинок_з_лапками)."""
    counter = Counter()
    total_cells = 0
    cells_with_newline = 0
    cells_with_quote = 0
    for table in tables:
        for row in table.rows:
            for cell in row:
                if cell is None:
                    continue
                s = str(cell)
                total_cells += 1
                counter.update(s)
                if "\n" in s or "\r" in s:
                    cells_with_newline += 1
                if '"' in s:
                    cells_with_quote += 1
    return counter, total_cells, cells_with_newline, cells_with_quote


def get_candidates_report(char_counter):
    """Кандидати на роздільник: безпечні (відсутні в даних) першими, за пріоритетом."""
    report = [
        {"char": ch, "name": name, "count": char_counter.get(ch, 0), "safe": char_counter.get(ch, 0) == 0}
        for ch, name, _ in DELIMITER_CANDIDATES
    ]
    report.sort(key=lambda r: (not r["safe"], _PRIORITY[r["char"]]))
    return report


# ------------------------------------------------------------------ числа

def normalize_number(s):
    """'3 600,00' -> '3600.00', '0,5' -> '0.5'. Усе інше (коди, текст, цілі з нулями) без змін."""
    t = s.strip()
    if _NUM_SPACED.fullmatch(t):
        t = re.sub(r"[   ]", "", t).replace(",", ".")
        return t
    if _NUM_DECIMAL_COMMA.fullmatch(t):
        return t.replace(",", ".")
    return s


# ------------------------------------------------------------------ підготовка рядків

def _pad(rows):
    width = max((len(r) for r in rows), default=0)
    return [list(r) + [None] * (width - len(r)) for r in rows]


def _fill_merges(rows, merges):
    """Об'єднання -> плоска таблиця: назви колонок шапки унікальні й непорожні
    (продовження об'єднаної комірки -> 'Назва_2'), об'єднане по вертикалі заповнюється значенням."""
    if not rows:
        return rows
    header = rows[0]
    for r0, c0, r1, c1 in merges:
        if r0 == 0 and c1 > c0:
            base = str(header[c0] or "").strip()
            for k, c in enumerate(range(c0 + 1, c1 + 1), start=2):
                if not str(header[c] or "").strip() and base:
                    header[c] = f"{base}_{k}"
        if r1 > r0 and r0 > 0:
            for r in range(r0 + 1, r1 + 1):
                for c in range(c0, c1 + 1):
                    if not str(rows[r][c] or "").strip():
                        rows[r][c] = rows[r0][c0]
    for c, v in enumerate(header):
        if not str(v or "").strip():
            header[c] = f"Колонка_{c + 1}"
    return rows


def _clean_value(v, flatten_newlines, sanitize_quotes, normalize_numbers, replace_char=None):
    if v is None:
        return ""
    if isinstance(v, float):
        s = str(int(v)) if v.is_integer() else repr(v)
    else:
        s = str(v)
    if flatten_newlines:
        s = s.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")
    if sanitize_quotes:
        s = s.replace('"', "'")
    if normalize_numbers:
        s = normalize_number(s)
    if replace_char:
        s = s.replace(replace_char, " " if replace_char == "," else ",")
    return s


def prepare_rows(
    rows,
    merges=(),
    *,
    trim_trailing_empty=True,
    fill_merges=False,
    normalize_numbers=False,
    flatten_newlines=True,
    sanitize_quotes=True,
    replace_char=None,
):
    """Готує рядки для CSV: рядки-списки рядків без None. replace_char - символ (роздільник),
    який замінюється в тексті на ',' (або пробіл), щоб поля не потребували лапок."""
    data = _pad(get_data_rows(rows, trim_trailing_empty))
    if fill_merges:
        data = _fill_merges(data, [m for m in merges if m[2] < len(data)])
    return [
        [_clean_value(v, flatten_newlines, sanitize_quotes, normalize_numbers, replace_char) for v in row]
        for row in data
    ]


def write_rows_csv(rows, delimiter, out_path, use_bom=True):
    """Запис уже підготовлених рядків (напр. у розкладці профілю)."""
    with open(out_path, "w", newline="", encoding="utf-8-sig" if use_bom else "utf-8") as f:
        csv.writer(f, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL).writerows(rows)


def convert_rows_to_csv(rows, delimiter, out_path, use_bom=True, merges=(), **options):
    encoding = "utf-8-sig" if use_bom else "utf-8"
    with open(out_path, "w", newline="", encoding=encoding) as f:
        writer = csv.writer(f, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
        writer.writerows(prepare_rows(rows, merges, **options))


def preview_csv(rows, delimiter, max_rows=5, merges=(), **options):
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=delimiter, quoting=csv.QUOTE_MINIMAL)
    writer.writerows(prepare_rows(rows, merges, **options)[:max_rows])
    return buf.getvalue()


# ------------------------------------------------------------------ Excel

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_name(name, used):
    base = _BAD_SHEET_CHARS.sub("_", name).strip("'")[:31] or "Таблиця"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def convert_tables_to_xlsx(tables, out_path):
    """Без очищення: значення як є. Кожна таблиця - окремий аркуш; об'єднані комірки зберігаються."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for table in tables:
        ws = wb.create_sheet(_sheet_name(table.name, used))
        widths = {}
        spanning = {(m[0] + 1, m[1] + 1) for m in table.merges if m[3] > m[1]}
        for r, row in enumerate(table.rows, start=1):
            for c, v in enumerate(row, start=1):
                if v is None or v == "":
                    continue
                if isinstance(v, str):
                    v = _ILLEGAL_XML.sub("", v)
                ws.cell(row=r, column=c, value=v)
                if (r, c) not in spanning:
                    widths[c] = max(widths.get(c, 0), min(len(str(v).split("\n")[0]), 60))
        for r0, c0, r1, c1 in table.merges:
            ws.merge_cells(start_row=r0 + 1, start_column=c0 + 1, end_row=r1 + 1, end_column=c1 + 1)
            in_table = not table.borders or any(
                b[0] <= r0 <= b[2] and b[1] <= c0 <= b[3] for b in table.borders
            )
            if in_table:
                ws.cell(row=r0 + 1, column=c0 + 1).alignment = Alignment(horizontal="center", vertical="center")
        for r0, c0, r1, c1 in table.borders:
            for r in range(r0 + 1, r1 + 2):
                for c in range(c0 + 1, c1 + 2):
                    ws.cell(row=r, column=c).border = box
        final = {c: max(8, w + 2) for c, w in widths.items()}
        for r0, c0, r1, c1 in table.merges:
            value = ws.cell(row=r0 + 1, column=c0 + 1).value
            in_table = not table.borders or any(
                b[0] <= r0 <= b[2] and b[1] <= c0 <= b[3] for b in table.borders
            )
            if c1 > c0 and in_table and isinstance(value, str):
                span = range(c0 + 1, c1 + 2)
                missing = min(len(value), 60) + 2 - sum(final.get(c, 8) for c in span)
                if missing > 0:
                    for c in span:
                        final[c] = final.get(c, 8) + missing / len(span)
        for c, w in final.items():
            ws.column_dimensions[get_column_letter(c)].width = w
    wb.save(out_path)


if __name__ == "__main__":
    import sys

    import sources

    path = sys.argv[1]
    tables = sources.load_pdf_tables(path) if sources.is_pdf(path) else sources.load_excel_tables(path)
    print("Таблиці:", [(t.name, t.n_rows, t.n_cols, t.merges) for t in tables])
    counter, total, with_nl, with_quote = analyze_characters(tables)
    print(f"Клітинок: {total}, унікальних символів: {len(counter)}, з переносом: {with_nl}, з лапками: {with_quote}")
    report = get_candidates_report(counter)
    safe = [r for r in report if r["safe"]]
    print("Безпечні роздільники:", [r["char"] for r in safe])
    if safe and tables:
        print(preview_csv(tables[0].rows, safe[0]["char"], max_rows=3, merges=tables[0].merges, fill_merges=True, normalize_numbers=True))
