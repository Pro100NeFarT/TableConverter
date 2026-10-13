"""Автотест PDF-потоку: генеруємо PDF (потрібен reportlab, лише для тестів) і перевіряємо розпізнавання."""
import csv
import io
import os
import tempfile

import openpyxl

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import converter
import sources

pdfmetrics.registerFont(TTFont("Arial", r"C:\Windows\Fonts\arial.ttf"))
body = ParagraphStyle("b", parent=getSampleStyleSheet()["Normal"], fontName="Arial", fontSize=10)

path = os.path.join(tempfile.gettempdir(), "tc_test.pdf")
doc = SimpleDocTemplate(path, pagesize=A4)

header = ["Артикул", "Найменування", "Кількість", "Ціна"]
data = [header]
for i in range(1, 61):
    name = f"Товар {i}"
    if i == 3:
        name = Paragraph('Сироватка "Ерл грей"<br/>другий рядок', body)
    data.append([f"НФ-{i:05d}", name, str(i * 2), f"{i * 10},50"])
t1 = Table(data, repeatRows=1)
t1.setStyle(
    TableStyle(
        [("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTNAME", (0, 0), (-1, -1), "Arial"), ("FONTSIZE", (0, 0), (-1, -1), 10)]
    )
)

plain = [["Назва", "Країна", "Сума"]] + [[f"Позиція {i}", "Україна", f"{i * 100}"] for i in range(1, 8)]
t2 = Table(plain, colWidths=[150, 150, 100])
t2.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "Arial"), ("FONTSIZE", (0, 0), (-1, -1), 10)]))

doc.build(
    [
        Paragraph("Накладна № 15 від 01.10.2026. Постачальник ТОВ Приклад.", body),
        Spacer(1, 12),
        t1,
        PageBreak(),
        Paragraph("Додаток. Таблиця без ліній:", body),
        Spacer(1, 12),
        t2,
    ]
)

# 1. Лінії + склеювання сторінок
tables = sources.load_pdf_tables(path, mode="lines", merge_continued=True)
print("lines+merge:", [(t.name, t.n_rows, t.n_cols) for t in tables])
assert len(tables) == 1, tables
t = tables[0]
assert t.name.startswith("Стор. 1-"), t.name
assert t.n_rows == 61 and t.n_cols == 4, (t.n_rows, t.n_cols)
assert t.rows[0] == header
assert sum(1 for r in t.rows if r == header) == 1, "повторена шапка має бути прибрана"
assert t.rows[60][0] == "НФ-00060"
assert "Шапку/абзац" not in str(t.rows) and not any("Накладна" in c for r in t.rows for c in r), "текст поза таблицею потрапив у таблицю"

# 2. Без склеювання
tables_nm = sources.load_pdf_tables(path, mode="lines", merge_continued=False)
print("lines no-merge:", [(t.name, t.n_rows) for t in tables_nm])
assert len(tables_nm) == 2 and tables_nm[0].n_rows + tables_nm[1].n_rows == 62

# 3. Авто: ще й таблиця без ліній на останній сторінці
tables_auto = sources.load_pdf_tables(path, mode="auto", merge_continued=True)
print("auto:", [(t.name, t.n_rows, t.n_cols) for t in tables_auto])
assert len(tables_auto) == 2
plain_t = tables_auto[1]
assert plain_t.n_cols == 3 and plain_t.rows[0] == ["Назва", "Країна", "Сума"], plain_t.rows[:2]
assert plain_t.n_rows == 8

tables_text = sources.load_pdf_tables(path, mode="text", merge_continued=True)
print("text:", [(t.name, t.n_rows, t.n_cols) for t in tables_text])
assert any(t.rows[0] == ["Назва", "Країна", "Сума"] and t.n_rows == 8 for t in tables_text)

# 4. Повний конвеєр: аналіз -> CSV
counter, total, with_nl, with_q = converter.analyze_characters([t])
print("клітинок:", total, "з переносом:", with_nl, "з лапками:", with_q)
report = converter.get_candidates_report(counter)
delim = [r for r in report if r["safe"]][0]["char"]
out = os.path.join(tempfile.gettempdir(), "tc_test.csv")
converter.convert_rows_to_csv(t.rows, delim, out)
with open(out, "r", encoding="utf-8-sig", newline="") as f:
    content = f.read()
rows = list(csv.reader(io.StringIO(content), delimiter=delim))
assert len(rows) == 61 and {len(r) for r in rows} == {4}
assert '"' not in content and content.count("\r\n") == 61
row3 = rows[3]
print("рядок з лапками і переносом:", row3)
assert "Ерл грей" in row3[1] and "другий рядок" in row3[1]
os.remove(out)

# 4б. Excel "1 в 1": текст поза таблицею + таблиця з рамками в одній сітці
sheets = sources.load_pdf_layout(path)
print("layout:", [(s.name, s.n_rows, s.n_cols) for s in sheets])
assert len(sheets) == 3
s1 = sheets[0]
assert "Накладна" in s1.rows[0][0], s1.rows[0]
hdr = next(i for i, r in enumerate(s1.rows) if "Артикул" in r)
assert hdr >= 1 and s1.borders, (hdr, len(s1.borders))
assert any(m[0] == 0 and m[3] > m[1] for m in s1.merges), "рядок тексту має бути розтягнутий на ширину"
xl1 = os.path.join(tempfile.gettempdir(), "tc_test_layout.xlsx")
converter.convert_tables_to_xlsx(sheets, xl1)
wb1 = openpyxl.load_workbook(xl1)
assert wb1.sheetnames == ["Стор. 1", "Стор. 2", "Стор. 3"]
w1 = wb1["Стор. 1"]
assert "Накладна" in w1["A1"].value and w1.cell(row=hdr + 1, column=1).border.left.style == "thin"
assert any("Артикул" == c.value for c in w1[hdr + 1])
os.remove(xl1)

# 5. Об'єднані комірки: шапка на 2 колонки + вертикальне об'єднання на 2-й сторінці, числа, Excel

path2 = os.path.join(tempfile.gettempdir(), "tc_test_merge.pdf")
inv = [["№", "Товар", "Кількість", "", "Ціна"]]
for i in range(1, 61):
    inv.append([str(i), f"Товар {i}", "1", "шт", f"{i * 1000 + 600:,}".replace(",", " ") + ",00"])
inv[46][0] = ""  # другий рядок вертикального об'єднання
t3 = Table(inv, repeatRows=1)
t3.setStyle(
    TableStyle(
        [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
            ("FONTNAME", (0, 0), (-1, -1), "Arial"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("SPAN", (2, 0), (3, 0)),
            ("SPAN", (0, 45), (0, 46)),
        ]
    )
)
SimpleDocTemplate(path2, pagesize=A4).build([t3])

tm = sources.load_pdf_tables(path2, mode="lines", merge_continued=True)
print("merge table:", [(t.name, t.n_rows, t.n_cols) for t in tm], tm[0].merges)
assert len(tm) == 1 and tm[0].n_rows == 61 and tm[0].n_cols == 5
assert (0, 2, 0, 3) in tm[0].merges, tm[0].merges
assert (45, 0, 46, 0) in tm[0].merges, tm[0].merges

prep = converter.prepare_rows(tm[0].rows, tm[0].merges, fill_merges=True, normalize_numbers=True)
assert prep[0] == ["№", "Товар", "Кількість", "Кількість_2", "Ціна"], prep[0]
assert prep[46][0] == prep[45][0] == "45", (prep[45][0], prep[46][0])
assert prep[1][4] == "1600.00", prep[1][4]
raw = converter.prepare_rows(tm[0].rows, tm[0].merges)
assert raw[0][3] == "" and raw[1][4] == "1 600,00"

xl = os.path.join(tempfile.gettempdir(), "tc_test_merge.xlsx")
converter.convert_tables_to_xlsx(tm, xl)
ws = openpyxl.load_workbook(xl).active
ranges = {str(r) for r in ws.merged_cells.ranges}
print("xlsx merges:", ranges, ws["E2"].value)
assert {"C1:D1", "A46:A47"} <= ranges, ranges
assert ws["E2"].value == "1 600,00" and ws["A2"].value == "1" and ws["D2"].value == "шт"  # без очищення
os.remove(xl)
os.remove(path2)

# 6. GUI-шлях: фонове читання PDF -> список таблиць -> аналіз -> прев'ю
import time

from app import App

app = App()
app.file_path = path
app.reload()
deadline = time.time() + 60
while app._loading and time.time() < deadline:
    app.update()
    time.sleep(0.05)
assert not app._loading, "читання PDF не завершилось"
assert [t.name for t in app.tables] == ["Стор. 1-2, табл. 1", "Стор. 3, табл. 1"], [t.name for t in app.tables]
assert len(app.tables_list.curselection()) == 2
app.analyze()
d = app.get_selected_delimiter()
assert d == ";", repr(d)
app.update()
preview = app.preview_text.get("1.0", "end")
assert "Артикул" in preview, preview
print("GUI: знайдено", len(app.tables), "таблиці, роздільник", repr(d), "прев'ю ок")
app.destroy()

os.remove(path)
print("ВСІ ПЕРЕВІРКИ PDF ПРОЙДЕНО")
