"""
Джерела таблиць: Excel (.xlsx/.xlsm) та PDF з текстовим шаром.
Кожне джерело повертає список Table(name, rows), далі працює спільний конвеєр.
"""
import os
from dataclasses import dataclass, field

import openpyxl

PDF_MODES = {
    "auto": "Авто (спершу за лініями, потім за текстом)",
    "lines": "Лише за лініями таблиці",
    "text": "За вирівнюванням тексту (таблиці без ліній)",
}


@dataclass
class Table:
    name: str
    rows: list
    merges: list = field(default_factory=list)  # (рядок0, колонка0, рядок1, колонка1), включно, з нуля
    borders: list = field(default_factory=list)  # прямокутники з рамкою, у тому ж форматі
    meta: list = field(default_factory=list)  # положення таблиць на аркуші "1 в 1" (для підказки налаштувань 1С)

    @property
    def n_rows(self):
        return len(self.rows)

    @property
    def n_cols(self):
        return max((len(r) for r in self.rows), default=0)


def is_pdf(path):
    return os.path.splitext(path)[1].lower() == ".pdf"


def load_excel_tables(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        return [
            Table(name, [list(r) for r in wb[name].iter_rows(values_only=True)])
            for name in wb.sheetnames
        ]
    finally:
        wb.close()


# --------------------------------------------------------------------------- PDF

_LINES = {"vertical_strategy": "lines", "horizontal_strategy": "lines"}


def _norm_cell(v):
    return " ".join(str(v).split()) if v is not None else ""


def _clean_pdf_rows(rows):
    """Нормалізує клітинки (None -> "", зайві пробіли) і відкидає повністю порожні рядки."""
    out = []
    for row in rows:
        cells = [_norm_cell(c) for c in row]
        if any(cells):
            out.append(cells)
    return out


def _detect_merges(t):
    """Об'єднані комірки таблиці pdfplumber: комірка, чия рамка охоплює кілька колонок/рядків."""
    xs = sorted({c[0] for c in t.cells})
    ys = [r.bbox[1] for r in t.rows]
    merges = []
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            if not cell:
                continue
            c1 = max((j for j in range(ci, len(xs)) if xs[j] < cell[2] - 1.0), default=ci)
            r1 = max((k for k in range(ri, len(ys)) if ys[k] < cell[3] - 1.0), default=ri)
            if c1 > ci or r1 > ri:
                merges.append((ri, ci, r1, c1))
    return merges


def _extract_lined(t):
    """Рядки таблиці (без порожніх) і об'єднання з перерахованими індексами рядків."""
    grid = t.extract()
    kept = [i for i, row in enumerate(grid) if any(_norm_cell(c) for c in row)]
    new_index = {old: new for new, old in enumerate(kept)}
    rows = [[_norm_cell(c) for c in grid[i]] for i in kept]
    merges = []
    for r0, c0, r1, c1 in _detect_merges(t):
        if r0 not in new_index:
            continue
        nr0 = new_index[r0]
        nr1 = max(new_index[i] for i in kept if r0 <= i <= r1)
        if nr1 > nr0 or c1 > c0:
            merges.append((nr0, c0, nr1, c1))
    return rows, merges


def _is_real_table(rows):
    if len(rows) < 2:
        return False
    if max(len(r) for r in rows) < 2:
        return False
    total = sum(len(r) for r in rows)
    filled = sum(1 for r in rows for c in r if c)
    return filled / total >= 0.3


def _line_tables(page):
    found = []
    for t in page.find_tables(table_settings=_LINES):
        rows, merges = _extract_lined(t)
        if _is_real_table(rows):
            found.append((rows, merges))
    return found


def _line_cells(page):
    """Слова сторінки -> [(top, [[x0, x1, текст], ...])]: рядки за вертикаллю, комірки за проміжками."""
    words = page.extract_words(keep_blank_chars=False, use_text_flow=False, x_tolerance=1.5, y_tolerance=2)
    if not words:
        return []
    heights = sorted(w["bottom"] - w["top"] for w in words)
    gap_threshold = max(6.0, 0.9 * heights[len(heights) // 2])

    words.sort(key=lambda w: (w["top"], w["x0"]))
    lines = []
    for w in words:
        if lines and abs(w["top"] - lines[-1][0]["top"]) <= 3:
            lines[-1].append(w)
        else:
            lines.append([w])

    result = []
    for ln in lines:
        ln.sort(key=lambda w: w["x0"])
        cells = [[ln[0]["x0"], ln[0]["x1"], ln[0]["text"]]]
        for w in ln[1:]:
            if w["x0"] - cells[-1][1] <= gap_threshold:
                cells[-1][1] = w["x1"]
                cells[-1][2] += " " + w["text"]
            else:
                cells.append([w["x0"], w["x1"], w["text"]])
        result.append((ln[0]["top"], cells))
    return result


def _text_tables(page):
    """Таблиці без ліній: слова групуються в рядки за вертикаллю, у комірки - за невеликими
    проміжками, колонки - за перекриттям комірок по горизонталі. Блок = послідовні рядки
    принаймні з двома комірками (мінімум 3 рядки)."""
    line_cells = [cells for _, cells in _line_cells(page)]
    if not line_cells:
        return []

    blocks, current = [], []
    for cells in line_cells:
        if len(cells) >= 2:
            current.append(cells)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)

    tables = []
    for block in blocks:
        if len(block) < 3:
            continue
        columns = []
        for x0, x1 in sorted((c[0], c[1]) for cells in block for c in cells):
            if columns and x0 <= columns[-1][1]:
                columns[-1][1] = max(columns[-1][1], x1)
            else:
                columns.append([x0, x1])
        rows = []
        for cells in block:
            row = [""] * len(columns)
            for x0, x1, text in cells:
                center = (x0 + x1) / 2
                for ci, (a, b) in enumerate(columns):
                    if a <= center <= b:
                        row[ci] = (row[ci] + " " + text).strip()
                        break
            rows.append(row)
        if _is_real_table(rows):
            tables.append(rows)
    return tables


def _page_tables(page, mode):
    """Список (метод, рядки, об'єднання) таблиць сторінки; auto - за лініями, інакше за текстом."""
    if mode in ("lines", "auto"):
        found = _line_tables(page)
        if found:
            return [("lines", rows, merges) for rows, merges in found]
    if mode in ("text", "auto"):
        return [("text", rows, []) for rows in _text_tables(page)]
    return []


def extract_pdf_tables(path, mode="auto", progress=None):
    """Повертає список (сторінка, номер_на_сторінці, рядки, метод, об'єднання) знайдених таблиць."""
    import pdfplumber

    result = []
    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        for page_no, page in enumerate(pdf.pages, 1):
            if progress:
                progress(page_no, total)
            for idx, (method, rows, merges) in enumerate(_page_tables(page, mode), 1):
                result.append((page_no, idx, rows, method, merges))
    return result


def _header_key(row):
    return [c.lower() for c in row]


def pdf_tables_to_tables(found, merge_continued=True):
    """Перетворює знайдене на Table; опційно склеює таблиці, що тривають на наступній сторінці
    (та сама кількість колонок, перша на сторінці, попередня — остання на своїй сторінці),
    прибираючи повторену шапку."""
    per_page = {}
    for page_no, idx, _, _, _ in found:
        per_page[page_no] = max(per_page.get(page_no, 0), idx)

    groups = []
    for page_no, idx, rows, method, merges in found:
        ncols = max(len(r) for r in rows)
        if merge_continued and groups:
            g = groups[-1]
            g_ncols = max(len(r) for r in g["rows"])
            if (
                idx == 1
                and page_no == g["last_page"] + 1
                and g["last_idx"] == per_page[g["last_page"]]
                and ncols == g_ncols
                and method == g["method"]
            ):
                drop = _header_key(rows[0]) == _header_key(g["rows"][0])
                offset = len(g["rows"]) - (1 if drop else 0)
                for r0, c0, r1, c1 in merges:
                    if drop and r0 == 0:
                        continue
                    g["merges"].append((r0 + offset, c0, r1 + offset, c1))
                g["rows"].extend(rows[1:] if drop else rows)
                g["last_page"] = page_no
                g["last_idx"] = idx
                continue
        groups.append(
            {
                "first_page": page_no,
                "last_page": page_no,
                "idx": idx,
                "last_idx": idx,
                "rows": list(rows),
                "merges": list(merges),
                "method": method,
            }
        )

    tables = []
    for g in groups:
        if g["first_page"] == g["last_page"]:
            pages = f"Стор. {g['first_page']}"
        else:
            pages = f"Стор. {g['first_page']}-{g['last_page']}"
        tables.append(Table(f"{pages}, табл. {g['idx']}", g["rows"], g["merges"]))
    return tables


def load_pdf_tables(path, mode="auto", merge_continued=True, progress=None):
    found = extract_pdf_tables(path, mode=mode, progress=progress)
    return pdf_tables_to_tables(found, merge_continued=merge_continued)


# ----------------------------------------------------------- PDF "1 в 1" (розкладка сторінки)

def _in_any(bboxes, x, y):
    return any(b[0] - 1 <= x <= b[2] + 1 and b[1] - 1 <= y <= b[3] + 1 for b in bboxes)


def _grid_starts(xs, tol=3.0):
    groups = []
    for x in sorted(xs):
        if groups and x - groups[-1][-1] <= tol:
            groups[-1].append(x)
        else:
            groups.append([x])
    return [sum(g) / len(g) for g in groups]


def _page_layout(page, number):
    """Сторінка як аркуш: рядки тексту поза таблицями + таблиці з рамками й об'єднаннями,
    усе в одній сітці колонок. Повертає Table або None для порожньої сторінки."""
    tables = []
    for t in page.find_tables(table_settings=_LINES):
        if _is_real_table(_extract_lined(t)[0]):
            tables.append(t)
    bboxes = [t.bbox for t in tables]
    body = page
    if bboxes:
        body = page.filter(
            lambda o: o.get("object_type") != "char"
            or not _in_any(bboxes, (o["x0"] + o["x1"]) / 2, (o["top"] + o["bottom"]) / 2)
        )

    entries = []  # (top, порядок, дані)
    for top, cells in _line_cells(body):
        entries.append((top, 0, ("text", cells)))
    for ti, t in enumerate(tables):
        grid = t.extract()
        for ri, row in enumerate(t.rows):
            entries.append((row.bbox[1], 1, ("trow", ti, ri, row.cells, grid[ri])))
    if not entries:
        return None
    entries.sort(key=lambda e: (e[0], e[1]))

    xs = []
    for _, _, data in entries:
        if data[0] == "text":
            xs.extend(c[0] for c in data[1])
        else:
            xs.extend(c[0] for c in data[3] if c)
    starts = _grid_starts(xs)
    ncols = len(starts)

    def col_of(x):
        return min(range(ncols), key=lambda i: abs(starts[i] - x))

    out_index = {}  # (таблиця, рядок таблиці) -> рядок аркуша
    for r, (_, _, data) in enumerate(entries):
        if data[0] == "trow":
            out_index[(data[1], data[2])] = r
    row_tops = {ti: [row.bbox[1] for row in t.rows] for ti, t in enumerate(tables)}

    rows = [[""] * ncols for _ in entries]
    merges, borders = [], []

    def put(r, c, text):
        rows[r][c] = (rows[r][c] + " " + text).strip() if rows[r][c] else text

    for r, (_, _, data) in enumerate(entries):
        if data[0] == "text":
            cells = data[1]
            for i, (x0, _x1, text) in enumerate(cells):
                c0 = col_of(x0)
                c1 = max(c0, col_of(cells[i + 1][0]) - 1) if i + 1 < len(cells) else ncols - 1
                put(r, c0, text)
                if c1 > c0 and not any(m[0] == r and m[1] == c0 for m in merges):
                    merges.append((r, c0, r, c1))
        else:
            _, ti, ri, cells, texts = data
            for ci, cell in enumerate(cells):
                if not cell:
                    continue
                c0 = col_of(cell[0])
                c1 = max((j for j in range(c0, ncols) if starts[j] < cell[2] - 2.0), default=c0)
                k = max((k for k in range(ri, len(row_tops[ti])) if row_tops[ti][k] < cell[3] - 1.0), default=ri)
                r1 = out_index[(ti, k)]
                put(r, c0, _norm_cell(texts[ci]))
                if c1 > c0 or r1 > r:
                    merges.append((r, c0, r1, c1))
                borders.append((r, c0, r1, c1))

    meta = []
    for ti, t in enumerate(tables):
        grid = t.extract()
        filled = [ri for ri, row in enumerate(grid) if any(_norm_cell(c) for c in row)]
        if not filled:
            continue
        table_xs = sorted({c[0] for c in t.cells})
        meta.append(
            {
                "header": [_norm_cell(c) for c in grid[filled[0]]],
                "header_row": out_index[(ti, filled[0])],
                "last_row": out_index[(ti, filled[-1])],
                "col_map": {ci: col_of(x) for ci, x in enumerate(table_xs)},
            }
        )
    return Table(f"Стор. {number}", rows, merges, borders, meta)


def load_pdf_layout(path, progress=None):
    """Усі сторінки PDF як аркуші 'один в один': текст, реквізити, таблиці, підписи."""
    import pdfplumber

    sheets = []
    with pdfplumber.open(path) as pdf:
        total = len(pdf.pages)
        for number, page in enumerate(pdf.pages, 1):
            if progress:
                progress(number, total)
            sheet = _page_layout(page, number)
            if sheet:
                sheets.append(sheet)
    return sheets
