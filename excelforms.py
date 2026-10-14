"""
Excel-файли (.xlsx/.xlsm/.xls): читання аркушів з об'єднаними комірками і розпізнавання таблиці
всередині друкованої форми 1С (заголовок і реквізити зверху, підсумки знизу, шапка з об'єднаними
комірками, дані в вузьких "технічних" колонках).
"""
import os
import re

import openpyxl

_BIG_FILE = 25 * 1024 * 1024

_GROUPS = {
    "name": r"товар|найменув|назва|наименов|номенклатур|послуг|description",
    "qty": r"к[іi]льк|к[іi]л-?сть|к-сть|кол-?во|количеств|\bqty\b|quantity",
    "price": r"ц[іi]на|цена|price",
    "sum": r"^сума|^сумма|вартість|стоимость|amount|^всього|^всего",
}
_TOTAL_RE = re.compile(r"^\s*(разом|всього|усього|итого|всего|total)\b", re.I)


def _norm(value):
    return " ".join(str(value).split()).lower() if value is not None else ""


def _is_numeric(value):
    return isinstance(value, (int, float)) or bool(re.fullmatch(r"[\d\s.,\-]+", str(value)))


def _split_vat_header(text):
    """'Ціна без / з ПДВ, грн.' -> ('Ціна без ПДВ, грн.', 'Ціна з ПДВ, грн.'); інакше друга колонка - '(2)'."""
    pattern = re.compile(r"без\s*/\s*з\s*(пдв|ндс)", re.I)
    if pattern.search(text):
        return pattern.sub(lambda m: f"без {m.group(1)}", text), pattern.sub(lambda m: f"з {m.group(1)}", text)
    return text, f"{text} (2)"


def _num(value):
    return int(value) if isinstance(value, float) and value.is_integer() else value


def read_grids(path):
    """[(назва_аркуша, значення[рядок][колонка], об'єднання [(r0, c0, r1, c1) включно])]."""
    if os.path.splitext(path)[1].lower() == ".xls":
        return _read_xls(path)
    return _read_xlsx(path)


def _read_xlsx(path):
    big = os.path.getsize(path) > _BIG_FILE
    wb = openpyxl.load_workbook(path, data_only=True, read_only=big)
    try:
        grids = []
        for ws in wb.worksheets:
            values = [list(row) for row in ws.iter_rows(values_only=True)]
            merges = [] if big else [
                (r.min_row - 1, r.min_col - 1, r.max_row - 1, r.max_col - 1) for r in ws.merged_cells.ranges
            ]
            grids.append((ws.title, values, merges))
        return grids
    finally:
        wb.close()


def _read_xls(path):
    import xlrd

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except xlrd.compdoc.CompDocError:  # файл, збережений сторонньою програмою, зі «зламаною» структурою контейнера
        book = xlrd.open_workbook(path, formatting_info=True, ignore_workbook_corruption=True)
    grids = []
    for sheet in book.sheets():
        values = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                kind, raw = sheet.cell_type(r, c), sheet.cell_value(r, c)
                if kind == xlrd.XL_CELL_NUMBER:
                    row.append(float(raw))
                elif kind == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(raw, book.datemode))
                elif kind == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(raw))
                elif kind == xlrd.XL_CELL_TEXT and str(raw).strip() != "":
                    row.append(raw)
                else:
                    row.append(None)
            values.append(row)
        merges = [(rlo, clo, rhi - 1, chi - 1) for rlo, rhi, clo, chi in sheet.merged_cells]
        grids.append((sheet.name, values, merges))
    return grids


def format_value(value):
    """Значення клітинки для тексту рядків: числа як у друкованій формі (2 656 500,00)."""
    if isinstance(value, float):
        return format(value, ",.2f").replace(",", " ").replace(".", ",")
    return str(value)


def sheet_lines(values):
    """Текст аркуша по рядках; комірки з переносами розбиваються на окремі рядки."""
    lines = []
    for row in values:
        parts = [format_value(v) for v in row if v is not None and str(v).strip() != ""]
        for line in "\t".join(parts).split("\n"):
            line = re.sub(r"[  ]{2,}", " ", line).strip()
            if line:
                lines.append(line)
    return lines


def _header_row(values):
    for r, row in enumerate(values):
        texts = [_norm(v) for v in row if v is not None and str(v).strip() != ""]
        found = {g for g, pattern in _GROUPS.items() if any(re.search(pattern, t) for t in texts)}
        if {"name", "qty"} <= found and found & {"price", "sum"}:
            return r
    return None


def detect_table(values, merges):
    """(рядки, об'єднання шапки) таблиці форми або None. Колонки - за шапкою; якщо під об'єднаною
    коміркою шапки лежать кілька колонок даних (кількість і одиниця), вони стають окремими колонками."""
    header_r = _header_row(values)
    if header_r is None:
        return None
    spans = {(r0, c0): (r1, c1) for r0, c0, r1, c1 in merges}

    header = []  # (колонка, текст, остання колонка)
    data_start = header_r + 1
    for c, value in enumerate(values[header_r]):
        if value is None or str(value).strip() == "":
            continue
        r1, c1 = spans.get((header_r, c), (header_r, c))
        header.append((c, str(value), c1))
        data_start = max(data_start, r1 + 1)

    def span_of(group):
        for c, text, c1 in header:
            if re.search(_GROUPS[group], _norm(text)):
                return c, c1
        return None

    name_span, qty_span = span_of("name"), span_of("qty")

    def filled(row, span):
        return any(row[c] is not None and str(row[c]).strip() != "" for c in range(span[0], min(span[1] + 1, len(row))))

    items, cont = [], {}  # cont: рядок товару -> рядок продовження (друга пара ціни/суми "з ПДВ")
    for r in range(data_start, len(values)):
        row = values[r]
        texts = [v for v in row if v is not None and str(v).strip() != ""]
        if not texts:
            if items:
                break
            continue
        if any(isinstance(v, str) and _TOTAL_RE.match(_norm(v)) for v in texts):
            break
        if not (filled(row, name_span) or filled(row, qty_span)):
            if items and r == items[-1] + 1 and items[-1] not in cont and all(_is_numeric(v) for v in texts):
                cont[items[-1]] = r
                continue
            if items:
                break
            continue
        items.append(r)
    if not items:
        return None
    if len(cont) < max(1, len(items) // 2):
        cont = {}

    def cell(r, c):
        return values[r][c] if 0 <= r < len(values) and c < len(values[r]) else None

    columns, head_row, head_merges = [], [], []  # columns: (колонка, з рядка продовження?)
    for c, text, c1 in header:
        anchors = sorted(
            {cc for r in items for cc in range(c, min(c1 + 1, len(values[r]))) if values[r][cc] is not None and str(values[r][cc]).strip() != ""}
        ) or [c]
        if len(anchors) == 1 and any(cell(cont[r], anchors[0]) not in (None, "") for r in cont):
            first, second = _split_vat_header(text)
            columns += [(anchors[0], False), (anchors[0], True)]
            head_row += [first, second]
            continue
        start = len(columns)
        for k, anchor in enumerate(anchors):
            columns.append((anchor, False))
            head_row.append(text if k == 0 else "")
        if len(anchors) > 1:
            head_merges.append((0, start, 0, start + len(anchors) - 1))

    rows = [head_row] + [
        [_num(cell(cont[r], c) if (second and r in cont) else cell(r, c)) if not (second and r not in cont) else None
         for c, second in columns]
        for r in items
    ]
    return rows, head_merges
